"""Export helpers: CSV, XLSX, PDF — clean output, no copyright clutter."""
import io, csv

def csv_bytes(headers, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow(list(r))
    return buf.getvalue().encode("utf-8-sig")


def xlsx_bytes(headers, rows, sheet="Data"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:30]

    ws.append(list(headers))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="059669")
        c.alignment = Alignment(vertical="center", horizontal="left")

    for r in rows:
        ws.append(list(r))

    for i, col in enumerate(ws.iter_cols(), start=1):
        letter = get_column_letter(i)
        max_len = 0
        for c in col:
            try:
                max_len = max(max_len, len(str(c.value or "")))
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(45, max(12, max_len + 2))

    ws.freeze_panes = "A2"

    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def pdf_bytes(title, headers, rows, subtitle=""):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer)
    from reportlab.lib.units import mm

    b = io.BytesIO()
    # Landscape A4 fits wide tables better; keeps "A4" as required.
    page = landscape(A4)
    doc = SimpleDocTemplate(
        b, pagesize=page,
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=title, author="",
    )

    st = getSampleStyleSheet()
    h_style = ParagraphStyle("H", parent=st["Title"],
                             fontSize=15, leading=18,
                             textColor=colors.HexColor("#0f172a"),
                             spaceAfter=4, alignment=0)
    s_style = ParagraphStyle("S", parent=st["Normal"],
                             fontSize=9, leading=12,
                             textColor=colors.HexColor("#64748b"),
                             spaceAfter=8)

    story = [Paragraph(title, h_style)]
    if subtitle:
        story.append(Paragraph(subtitle, s_style))
    story.append(Spacer(1, 4))

    data = [list(headers)] + [list(map(str, r)) for r in rows]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND",     (0, 0), (-1, 0), colors.HexColor("#059669")),
        ("TEXTCOLOR",      (0, 0), (-1, 0), colors.white),
        ("FONTNAME",       (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",       (0, 0), (-1, -1), 8),
        ("GRID",           (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f8fafc")]),
        ("VALIGN",         (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING",    (0, 0), (-1, -1), 4),
        ("RIGHTPADDING",   (0, 0), (-1, -1), 4),
        ("TOPPADDING",     (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING",  (0, 0), (-1, -1), 3),
    ]))
    story.append(t)
    doc.build(story)
    return b.getvalue()
